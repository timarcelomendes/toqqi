"""Leitura de planilhas (.csv, .xlsx, .xls) para linhas de texto."""
import csv
import io
from datetime import date, datetime

from toqqi.core.errors import AppError
from toqqi.core.texto import normalizar_cabecalho

MAX_BYTES = 5 * 1024 * 1024
MAX_LINHAS = 20_000
EXTENSOES = (".csv", ".xlsx", ".xls")


def erro_arquivo(msg: str) -> AppError:
    return AppError(422, "arquivo_invalido", msg, {"arquivo": msg})


def _celula(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "sim" if v else "não"
    if isinstance(v, datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v).strip()


def _decodificar(conteudo: bytes) -> str:
    for cod in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return conteudo.decode(cod)
        except UnicodeDecodeError:
            continue
    raise erro_arquivo("Não conseguimos ler o texto do arquivo.")


def _ler_csv(conteudo: bytes):
    texto = _decodificar(conteudo)
    primeira = next((ln for ln in texto.splitlines() if ln.strip()), "")
    delimitador = max((";", ",", "\t"), key=primeira.count) if primeira else ";"
    for i, linha in enumerate(csv.reader(io.StringIO(texto, newline=""), delimiter=delimitador), start=1):
        yield i, [_celula(c) for c in linha]


def _ler_xlsx(conteudo: bytes):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - arquivo corrompido ou não é xlsx
        raise erro_arquivo("Não conseguimos abrir a planilha. Salve de novo como .xlsx ou .csv.")
    try:
        ws = wb.worksheets[0]
        for i, linha in enumerate(ws.iter_rows(values_only=True), start=1):
            yield i, [_celula(c) for c in linha]
    finally:
        wb.close()


def _ler_xls(conteudo: bytes):
    import xlrd

    try:
        wb = xlrd.open_workbook(file_contents=conteudo)
    except Exception:  # noqa: BLE001
        raise erro_arquivo("Não conseguimos abrir a planilha. Salve de novo como .xlsx ou .csv.")
    sh = wb.sheet_by_index(0)
    for r in range(sh.nrows):
        valores = []
        for c in range(sh.ncols):
            cel = sh.cell(r, c)
            if cel.ctype == xlrd.XL_CELL_DATE:
                valores.append(_celula(xlrd.xldate_as_datetime(cel.value, wb.datemode)))
            elif cel.ctype == xlrd.XL_CELL_BOOLEAN:
                valores.append(_celula(bool(cel.value)))
            else:
                valores.append(_celula(cel.value))
        yield r + 1, valores


def ler(nome: str, conteudo: bytes) -> tuple[list[str], list[list]]:
    """Devolve (colunas, linhas) com linhas = [[nº da linha no arquivo, valor1, valor2, ...], ...]."""
    ext = "." + nome.rsplit(".", 1)[-1].lower() if "." in nome else ""
    if ext not in EXTENSOES:
        raise erro_arquivo("Envie um arquivo .csv, .xlsx ou .xls.")
    if len(conteudo) > MAX_BYTES:
        raise erro_arquivo("O arquivo passa de 5 MB. Divida a planilha em partes menores.")
    if not conteudo:
        raise erro_arquivo("O arquivo está vazio.")
    leitor = {".csv": _ler_csv, ".xlsx": _ler_xlsx, ".xls": _ler_xls}[ext](conteudo)

    colunas: list[str] | None = None
    linhas: list[list] = []
    for numero, valores in leitor:
        if not any(valores):
            continue
        if colunas is None:
            colunas = []
            while valores and not valores[-1]:  # tira colunas sem nome do fim
                valores = valores[:-1]
            for i, c in enumerate(valores):
                nome_col = c or f"coluna_{i + 1}"
                base, n = nome_col, 2
                while nome_col in colunas:
                    nome_col = f"{base} ({n})"
                    n += 1
                colunas.append(nome_col)
            continue
        if len(linhas) >= MAX_LINHAS:
            raise erro_arquivo("A planilha passa de 20.000 linhas. Divida em partes menores.")
        valores = (valores + [""] * len(colunas))[:len(colunas)]
        linhas.append([numero, *valores])
    if not colunas:
        raise erro_arquivo("Não encontramos o cabeçalho da planilha (a primeira linha com os nomes das colunas).")
    if not linhas:
        raise erro_arquivo("A planilha não tem nenhuma linha de dados.")
    return colunas, linhas


# ---- campos e nomes equivalentes -------------------------------------------

CAMPOS = [
    ("nome", "Nome do contato", True),
    ("email", "E-mail", False),
    ("telefone", "Telefone / WhatsApp", False),
    ("empresa", "Empresa", False),
    ("documento_empresa", "CNPJ/CPF da empresa", False),
    ("cargo", "Cargo", False),
    ("perfil", "Perfil", False),
    ("grupo", "Grupo", False),
    ("segmento", "Segmento", False),
    ("responsavel", "Responsável", False),
    ("valor_mensal", "Valor mensal (R$)", False),
    ("cliente_desde", "Cliente desde", False),
    ("codigo_externo", "Código no ERP", False),
    ("ativo", "Ativo", False),
]
CHAVES_CAMPOS = [c for c, _, _ in CAMPOS]
ROTULOS = {c: r for c, r, _ in CAMPOS}

_ALIASES = {
    "nome": ["nome", "nome_contato", "nome_do_contato", "contato", "nome_completo", "pessoa"],
    "email": ["email", "e_mail", "email_cliente", "email_contato", "email_do_contato", "e_mail_cliente",
              "e_mail_contato", "correio_eletronico", "mail"],
    "telefone": ["telefone", "whatsapp", "whats", "zap", "celular", "fone", "tel", "telefone_celular",
                 "numero_whatsapp", "telefone_whatsapp", "movel"],
    "empresa": ["empresa", "razao_social", "cliente", "nome_empresa", "nome_da_empresa", "nome_fantasia",
                "companhia", "empresa_cliente"],
    "documento_empresa": ["documento_empresa", "cnpj", "cpf", "cnpj_cpf", "cpf_cnpj", "documento",
                          "cnpj_empresa", "cnpj_da_empresa"],
    "cargo": ["cargo", "funcao"],
    "perfil": ["perfil", "perfil_contato", "perfil_do_contato", "papel"],
    "grupo": ["grupo", "rede", "grupo_economico"],
    "segmento": ["segmento", "ramo", "setor", "ramo_de_atividade"],
    "responsavel": ["responsavel", "vendedor", "representante", "gerente_de_contas", "consultor", "executivo",
                    "executivo_de_contas", "carteira"],
    "valor_mensal": ["valor_mensal", "faturamento_mensal", "mensalidade", "valor", "ticket_medio", "mrr"],
    "cliente_desde": ["cliente_desde", "data_inicio", "inicio", "desde", "data_de_inicio", "data_cadastro"],
    "codigo_externo": ["codigo_externo", "codigo", "cod", "id_erp", "codigo_erp", "codigo_cliente", "cod_cliente",
                       "id_externo", "id_cliente"],
    "ativo": ["ativo", "situacao", "status", "ativa"],
}
ALIASES = {alias: campo for campo, lista in _ALIASES.items() for alias in lista}

# ---- respostas antigas ------------------------------------------------------

CAMPOS_RESPOSTAS = [
    ("email", "E-mail do contato", True),
    ("data", "Data da resposta", True),
    ("nota", "Nota (0 a 10)", True),
    ("empresa", "Empresa", False),
    ("comentario", "Comentário", False),
]
CHAVES_RESPOSTAS = [c for c, _, _ in CAMPOS_RESPOSTAS]
ROTULOS_RESPOSTAS = {c: r for c, r, _ in CAMPOS_RESPOSTAS}
OBRIGATORIOS_RESPOSTAS = [c for c, _, o in CAMPOS_RESPOSTAS if o]
_ALIASES_RESPOSTAS = {
    "email": ["email", "e_mail", "email_cliente", "e_mail_cliente", "email_contato", "e_mail_contato",
              "email_do_contato", "correio_eletronico", "mail"],
    "data": ["data", "data_resposta", "data_da_resposta", "dt", "dt_resposta", "respondida_em", "respondido_em"],
    "nota": ["nota", "nps", "nota_nps", "score", "pontuacao", "nota_do_cliente"],
    "empresa": ["empresa", "razao_social", "cliente", "nome_empresa", "nome_da_empresa", "nome_fantasia",
                "empresa_cliente"],
    "comentario": ["comentario", "comentarios", "motivo", "observacao", "observacoes", "obs", "justificativa",
                   "comentario_do_cliente"],
}
ALIASES_RESPOSTAS = {alias: campo for campo, lista in _ALIASES_RESPOSTAS.items() for alias in lista}


def sugerir_mapeamento(colunas: list[str], tipo: str = "contatos") -> dict[str, str | None]:
    aliases = ALIASES_RESPOSTAS if tipo == "respostas" else ALIASES
    usados: set[str] = set()
    saida: dict[str, str | None] = {}
    for c in colunas:
        campo = aliases.get(normalizar_cabecalho(c))
        if campo in usados:
            campo = None
        if campo:
            usados.add(campo)
        saida[c] = campo
    return saida
